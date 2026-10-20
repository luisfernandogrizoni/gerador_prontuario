"""
Motor dos documentos do sistema: prontuário, declarações e relatórios.

Cada documento é descrito UMA vez, como uma lista de blocos (título, parágrafo, tabela, formulário, assinatura…),
e este módulo o transforma em duas coisas:

  * um arquivo Word (.docx), com o timbre montado aqui mesmo — não há arquivo de modelo guardado em pasta;
  * uma página HTML (templates/documento.html) para ver e imprimir no navegador, com o mesmo timbre.

O texto do timbre mora em CABECALHO e RODAPE (abaixo) e as duas logos em static/timbre/: é o único lugar a mexer
quando o endereço, o CNPJ ou a logo mudarem. O texto de cada documento mora no módulo dele (prontuario.py,
declaracoes.py, relatorios.py).

Uma página nova começa onde um parágrafo tem `quebra_antes`; no Word isso é uma quebra de página e no navegador
cada página vira uma folha A4 com o seu timbre. Uma página que precisa ser diferente (em paisagem, ou sem o timbre,
como a capa do envelope) começa com um bloco `Pagina(...)`: no Word vira uma seção nova.
"""

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import ClassVar

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips

PASTA_ESTATICOS = Path(__file__).parent / "static"
PASTA_TIMBRE = PASTA_ESTATICOS / "timbre"
LOGO_TOPO = PASTA_TIMBRE / "logo-topo.jpg"
LOGO_RODAPE = PASTA_TIMBRE / "logo-rodape.jpg"

# quem emite (e assina) as declarações e os relatórios
EMISSOR_NOME = "Luís Fernando Grizoni Elias"
EMISSOR_CARGO = "Assistente Administrativo"
CIDADE = "Assis"

# ----------------------------------------------------------------- timbre
# Cada linha é uma lista de trechos (texto, estilo); estilo: "b" negrito, "vermelho", "link", "serifa".
CABECALHO = [
    [("Comunidade Terapêutica", "b")],
    [("Casa de Acolhida Restauração", "b")],
    [("Por uma só vida já valeu a pena!", "b vermelho")],
    [("Rodovia Raposo Tavares, km 435 – Assis – SP, Caixa Postal: 818", "")],
]
RODAPE = [
    [("Associação Restauração - ", "b serifa"), ("Nossa Missão é Amar", "b vermelho")],
    [("Escritório: Rua Sebastião da Silva Leite 1145 – Centro – Assis – São Paulo.", "")],
    [("Telefone: (18) 3323-4778 (18) 99618-1905 CEP 19814-371", "")],
    [("CNPJ 03.508.198/0001-07", "")],
    [("UTILIDADE PÚBLICA MUNICIPAL: Lei nº 3960/2000", "")],
    [("UTILIDADE PÚBLICA ESTADUAL: Lei nº 17591/2022", "")],
    [("www.restauração.org.br", "link"), ("    ", ""), ("restauração@restauração.org.br", "link")],
]
TAMANHO_CABECALHO, TAMANHO_RODAPE = 10, 9

# página A4 (em twips: 1 cm = 567). Em cima e embaixo a margem é a do prontuário; dos lados cada documento escolhe.
LARGURA_PAGINA, ALTURA_PAGINA = 11906, 16838
MARGEM_VERTICAL, MARGEM_LATERAL_PADRAO, DISTANCIA_TIMBRE = 720, 1134, 708


# ----------------------------------------------------------------- blocos
@dataclass
class Texto:
    """Um trecho de parágrafo com formatação própria."""
    texto: str
    negrito: bool = False
    italico: bool = False
    sublinhado: bool = False


def N(texto):
    """Atalho: trecho em negrito."""
    return Texto(texto, negrito=True)


def S(texto):
    """Atalho: trecho sublinhado."""
    return Texto(texto, sublinhado=True)


@dataclass
class Titulo:
    tipo: ClassVar[str] = "titulo"
    texto: str
    tamanho: float = 15
    antes: float = 6
    depois: float = 2


@dataclass
class Paragrafo:
    tipo: ClassVar[str] = "paragrafo"
    partes: list                       # str ou Texto
    alinhamento: str = "esquerda"      # esquerda | centro | direita | justificado
    antes: float = 0                   # pontos
    depois: float = 0
    primeira_linha: int = 0            # recuo da primeira linha, em twips
    entrelinha: float = 1.0
    tamanho: float | None = None       # None = tamanho do documento
    unido_ao_proximo: bool = False
    recuo: int = 0                     # recuo da esquerda, em twips
    suspenso: int = 0                  # recuo "pendurado" (usado nos itens numerados), em twips
    numero: str = ""                   # "1." — marca do item de uma lista numerada
    quebra_antes: bool = False         # começa uma página nova
    fonte: str | None = None           # None = fonte do documento

    def trechos(self):
        return [p if isinstance(p, Texto) else Texto(p) for p in self.partes]


@dataclass
class Coluna:
    titulo: str
    largura: int                       # twips
    alinhamento: str = "esquerda"


@dataclass
class Tabela:
    tipo: ClassVar[str] = "tabela"
    colunas: list
    linhas: list                       # lista de listas de texto
    negrito_linhas: tuple = ()         # posições (a partir de 0) das linhas em negrito
    antes: float = 0
    tamanho: float = 10
    cabecalho: bool = True             # a primeira linha da tabela é o título das colunas
    sombreado: bool = True             # título das colunas com fundo cinza
    cor_borda: str = "808080"
    altura_cabecalho: int = 0          # altura mínima da linha do título, em twips
    sem_bordas: bool = False

    def largura_total(self):
        return sum(c.largura for c in self.colunas)

    def porcentagens(self):
        total = self.largura_total()
        return [round(c.largura * 100 / total, 2) for c in self.colunas]


@dataclass
class Celula:
    texto: str = ""
    span: int = 1                      # quantas colunas da grade a célula ocupa
    negrito: bool = False


@dataclass
class Formulario:
    """Tabela de formulário: grade de colunas fina e células que ocupam várias colunas (a ficha de controle)."""
    tipo: ClassVar[str] = "formulario"
    grade: list                        # largura de cada coluna da grade, em twips
    linhas: list                       # lista de listas de Celula (a soma dos spans de cada linha é len(grade))
    alturas: list                      # altura mínima de cada linha, em twips
    tamanho: float = 12
    antes: float = 0

    def largura_total(self):
        return sum(self.grade)


@dataclass
class Pauta:
    """Linhas horizontais para escrever à mão."""
    tipo: ClassVar[str] = "pauta"
    linhas: int
    altura: float = 19                 # pontos entre uma linha e outra
    antes: float = 0


@dataclass
class Assinaturas:
    """Linhas de assinatura lado a lado, cada uma com o seu rótulo embaixo."""
    tipo: ClassVar[str] = "assinaturas"
    rotulos: list
    comprimento: int = 32              # tamanho de cada traço, em sublinhados
    tamanho_rotulo: float = 10
    antes: float = 0


@dataclass
class Assinatura:
    """Linha de assinatura com nome, cargo e o papel de quem assina."""
    tipo: ClassVar[str] = "assinatura"
    nome: str = EMISSOR_NOME
    cargo: str = EMISSOR_CARGO
    papel: str = "Responsável pela emissão"
    antes: float = 40
    registro: str = ""                 # registro profissional (CRESS...): sai embaixo do cargo sempre que a pessoa tem


@dataclass
class Pagina:
    """Começa uma página nova que não é a comum: em paisagem e/ou sem o timbre (a capa do envelope). Vale até o
    próximo `Pagina`; as páginas começadas por `quebra_antes` seguem o jeito da anterior."""
    tipo: ClassVar[str] = "pagina"
    paisagem: bool = False
    timbre: bool = True


@dataclass
class Imagem:
    """Figura de static/ (ex.: "timbre/logo-topo.jpg"), com a largura em milímetros."""
    tipo: ClassVar[str] = "imagem"
    arquivo: str
    largura: float
    alinhamento: str = "esquerda"      # esquerda | centro | direita
    antes: float = 0

    @property
    def caminho(self):
        return PASTA_ESTATICOS / self.arquivo


@dataclass
class Colunas:
    """Parágrafos lado a lado (a lista do enxoval). Cada coluna é uma lista de `Paragrafo`."""
    tipo: ClassVar[str] = "colunas"
    colunas: list
    espaco: int = 360                  # distância entre as colunas, em twips
    antes: float = 0


@dataclass
class Folha:
    """Uma página do documento já separada: os blocos dela e como ela é."""
    blocos: list = field(default_factory=list)
    paisagem: bool = False
    timbre: bool = True
    secao: bool = False                # começou por um `Pagina`: no Word é uma seção nova


@dataclass
class Documento:
    nome_arquivo: str                  # sem a extensão
    titulo: str                        # aparece na aba do navegador e nas propriedades do Word
    blocos: list
    fonte: str = "Arial"
    tamanho: float = 11
    voltar: str = ""                   # endereço da página de onde o documento foi aberto
    margem_lateral: int = MARGEM_LATERAL_PADRAO

    @property
    def largura_util(self):
        return LARGURA_PAGINA - 2 * self.margem_lateral

    def largura_da(self, folha):
        """Largura útil (entre as margens) de uma `Folha`, em twips: a paisagem é mais larga."""
        return (ALTURA_PAGINA if folha.paisagem else LARGURA_PAGINA) - 2 * self.margem_lateral

    @property
    def tem_tabela_larga(self):
        """Tabela com 4 colunas ou mais: no celular a folha rola de lado em vez de espremer as colunas."""
        return any(getattr(bloco, "tipo", "") == "tabela" and len(bloco.colunas) >= 4 for bloco in self.blocos)

    def paginas(self):
        """Os blocos separados em páginas (`Folha`): uma nova a cada `Pagina` e a cada bloco com quebra_antes."""
        folhas = []
        for bloco in self.blocos:
            if getattr(bloco, "tipo", "") == "pagina":
                folhas.append(Folha([], bloco.paisagem, bloco.timbre, secao=bool(folhas)))
                continue
            if not folhas or (getattr(bloco, "quebra_antes", False) and folhas[-1].blocos):
                anterior = folhas[-1] if folhas else None
                folhas.append(Folha([], anterior.paisagem if anterior else False, anterior.timbre if anterior else True))
            folhas[-1].blocos.append(bloco)
        return folhas


ALINHAMENTO_DOCX = {"esquerda": WD_ALIGN_PARAGRAPH.LEFT, "centro": WD_ALIGN_PARAGRAPH.CENTER,
                    "direita": WD_ALIGN_PARAGRAPH.RIGHT, "justificado": WD_ALIGN_PARAGRAPH.JUSTIFY}


def data_por_extenso(d):
    meses = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro",
             "outubro", "novembro", "dezembro"]
    return f"{d.day:02d} de {meses[d.month - 1]} de {d.year}"


@dataclass(frozen=True)
class Emissor:
    """Quem assina como responsável pela emissão (vem do cadastro de funcionários)."""
    nome: str = EMISSOR_NOME
    cargo: str = EMISSOR_CARGO
    registro: str = ""                 # registro profissional (CRESS da assistente social): assina junto do nome


def fecho(hoje, emissor=None):
    """'Assis, 09 de outubro de 2026.' alinhado à direita, seguido da assinatura de quem emite
    (`emissor`; sem ele, o de fábrica)."""
    emissor = emissor or Emissor()
    return [Paragrafo([f"{CIDADE}, {data_por_extenso(hoje)}."], "direita", antes=26, unido_ao_proximo=True),
            Assinatura(nome=emissor.nome, cargo=emissor.cargo, registro=emissor.registro)]


# ============================================================ gerador do Word
def _quebras(texto):
    """O navegador envia o Enter de um campo de texto como CR+LF, e o Word conta o CR e o LF como duas quebras: cada Enter virava uma linha em
    branco no documento. Aqui toda quebra vira uma só (LF)."""
    return (texto or "").replace("\r\n", "\n").replace("\r", "\n")


def _fonte(run, nome, tamanho, negrito=False, italico=False, cor=None, sublinhado=False):
    run.bold, run.italic, run.underline = negrito, italico, sublinhado
    run.font.size = Pt(tamanho)
    run.font.name = nome
    if cor:
        run.font.color.rgb = RGBColor.from_string(cor)
    fontes = run._element.get_or_add_rPr().get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fontes.set(qn(atributo), nome)


def _filho(elemento, nome, **atributos):
    novo = OxmlElement(nome)
    for chave, valor in atributos.items():
        novo.set(qn("w:" + chave), str(valor))
    elemento.append(novo)
    return novo


# ordem em que o Word exige os elementos de propriedades de tabela (fora dela, ele pode recusar o arquivo)
ORDEM_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
               "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def _propriedade_de_tabela(props, nome, **atributos):
    """Coloca o elemento em w:tblPr na posição certa e devolve-o (substitui um igual que já exista)."""
    for velho in props.findall(qn("w:" + nome)):
        props.remove(velho)
    novo = OxmlElement("w:" + nome)
    for chave, valor in atributos.items():
        novo.set(qn("w:" + chave), str(valor))
    seguintes = ORDEM_TBLPR[ORDEM_TBLPR.index(nome) + 1:]
    posterior = next((f for f in props if f.tag in {qn("w:" + n) for n in seguintes}), None)
    if posterior is not None:
        posterior.addprevious(novo)
    else:
        props.append(novo)
    return novo


def _configurar_tabela(tabela, largura, bordas=None, margens=(50, 90, 50, 90)):
    """Largura fixa, bordas (lados -> cor) e margens internas das células (topo, esquerda, base, direita)."""
    props = tabela._tbl.tblPr
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    _propriedade_de_tabela(props, "tblW", w=largura, type="dxa")
    quadro = _propriedade_de_tabela(props, "tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        cor = (bordas or {}).get(lado)
        if cor:
            _filho(quadro, "w:" + lado, val="single", sz=4, space=0, color=cor)
        else:
            _filho(quadro, "w:" + lado, val="nil")
    _propriedade_de_tabela(props, "tblLayout", type="fixed")
    cantos = _propriedade_de_tabela(props, "tblCellMar")
    for lado, valor in zip(("top", "left", "bottom", "right"), margens):
        _filho(cantos, "w:" + lado, w=valor, type="dxa")


def _grade(tabela, larguras):
    """Largura de cada coluna na grade da tabela e em cada célula (o Word usa as duas)."""
    for coluna, largura in zip(tabela.columns, larguras):
        coluna.width = Twips(largura)
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras):
            celula.width = Twips(largura)


def _linhas_do_timbre(celula, linhas, tamanho):
    """Escreve as linhas do timbre (centralizadas) numa célula do cabeçalho ou do rodapé."""
    primeiro = True
    for linha in linhas:
        p = celula.paragraphs[0] if primeiro else celula.add_paragraph()
        primeiro = False
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        for texto, estilo in linha:
            estilos = estilo.split()
            _fonte(p.add_run(texto), "Times New Roman" if "serifa" in estilos else "Arial", tamanho, negrito="b" in estilos,
                   cor="FF0000" if "vermelho" in estilos else "0563C1" if "link" in estilos else None, sublinhado="link" in estilos)


def _faixa_do_timbre(parte, logo, largura_logo_cm, linhas, tamanho, largura_util):
    """Cabeçalho/rodapé: logo à esquerda, texto centralizado na página (a terceira coluna só equilibra o espaço)."""
    parte.is_linked_to_previous = False
    lateral = int(largura_logo_cm * 567) + 120
    tabela = parte.add_table(rows=1, cols=3, width=Twips(largura_util))
    _configurar_tabela(tabela, largura_util, margens=(0, 0, 0, 0))
    _grade(tabela, (lateral, largura_util - 2 * lateral, lateral))
    esquerda, centro, direita = tabela.rows[0].cells
    for celula in (esquerda, centro, direita):
        celula.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = esquerda.paragraphs[0]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(logo), width=Twips(int(largura_logo_cm * 567)))
    _linhas_do_timbre(centro, linhas, tamanho)
    # o Word exige um parágrafo depois da tabela: o que já existia vira esse fecho, pequeno
    fecho_ = parte.paragraphs[0]
    fecho_._p.addprevious(tabela._tbl)
    fecho_.paragraph_format.space_before = fecho_.paragraph_format.space_after = Pt(0)
    _fonte(fecho_.add_run(""), "Arial", 2)
    fecho_.paragraph_format.line_spacing = Pt(2)


def _base_docx(doc):
    """Documento em branco com a fonte do documento como padrão (as páginas, com o timbre, vêm em _configurar_secao)."""
    d = DocxDocument()
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = doc.fonte, Pt(doc.tamanho)
    fontes = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fontes.set(qn(atributo), doc.fonte)
    return d


def _configurar_secao(secao, doc, folha):
    """Tamanho e direção do papel, margens e timbre (ou nada, nas páginas sem timbre) de uma seção."""
    largura, altura = (ALTURA_PAGINA, LARGURA_PAGINA) if folha.paisagem else (LARGURA_PAGINA, ALTURA_PAGINA)
    secao.orientation = WD_ORIENT.LANDSCAPE if folha.paisagem else WD_ORIENT.PORTRAIT
    secao.page_width, secao.page_height = Twips(largura), Twips(altura)
    secao.top_margin = secao.bottom_margin = Twips(MARGEM_VERTICAL)
    secao.left_margin = secao.right_margin = Twips(doc.margem_lateral)
    secao.header_distance = secao.footer_distance = Twips(DISTANCIA_TIMBRE)
    secao.header.is_linked_to_previous = secao.footer.is_linked_to_previous = False
    if folha.timbre:
        util = doc.largura_da(folha)
        _faixa_do_timbre(secao.header, LOGO_TOPO, 2.3, CABECALHO, TAMANHO_CABECALHO, util)
        _faixa_do_timbre(secao.footer, LOGO_RODAPE, 2.5, RODAPE, TAMANHO_RODAPE, util)
    else:                                                            # sem timbre: cabeçalho e rodapé vazios e pequenos
        for parte in (secao.header, secao.footer):
            for p in parte.paragraphs:
                _miniatura(p)


def _miniatura(paragrafo):
    """Deixa um parágrafo vazio com 1 pt de altura (a marca de fim de seção não pode empurrar uma página em branco)."""
    formato = paragrafo.paragraph_format
    formato.space_before = formato.space_after = Pt(0)
    formato.line_spacing = Pt(1)
    propriedades = paragrafo._p.get_or_add_pPr()
    marca = propriedades.find(qn("w:rPr"))
    if marca is None:                                                # o Word exige o rPr antes do sectPr
        marca = OxmlElement("w:rPr")
        posterior = next((e for e in (propriedades.find(qn("w:sectPr")), propriedades.find(qn("w:pPrChange"))) if e is not None), None)
        if posterior is not None:
            posterior.addprevious(marca)
        else:
            propriedades.append(marca)
    _filho(marca, "w:sz", val=2)


def _paragrafo_docx(d, doc, bloco, alvo=None, quebra=None):
    """`alvo` = onde escrever (o documento, ou uma célula de tabela); `quebra` troca bloco.quebra_antes."""
    p = (alvo or d).add_paragraph()
    p.alignment = ALINHAMENTO_DOCX[bloco.alinhamento]
    formato = p.paragraph_format
    formato.space_before, formato.space_after = Pt(bloco.antes), Pt(bloco.depois)
    formato.line_spacing = bloco.entrelinha
    formato.keep_with_next = bloco.unido_ao_proximo
    if bloco.quebra_antes if quebra is None else quebra:
        formato.page_break_before = True
    if bloco.recuo:
        formato.left_indent = Twips(bloco.recuo)
    if bloco.suspenso:
        formato.first_line_indent = Twips(-bloco.suspenso)
        formato.tab_stops.add_tab_stop(Twips(bloco.recuo))
    elif bloco.primeira_linha:
        formato.first_line_indent = Twips(bloco.primeira_linha)
    fonte, tamanho = bloco.fonte or doc.fonte, bloco.tamanho or doc.tamanho
    if bloco.numero:
        _fonte(p.add_run(bloco.numero + "\t"), fonte, tamanho)
    for t in bloco.trechos():
        _fonte(p.add_run(_quebras(t.texto)), fonte, tamanho, t.negrito, t.italico, sublinhado=t.sublinhado)
    return p


def _tabela_docx(d, doc, bloco):
    tabela = d.add_table(rows=1, cols=len(bloco.colunas))
    bordas = None if bloco.sem_bordas else {lado: bloco.cor_borda for lado in ("top", "left", "bottom", "right", "insideH", "insideV")}
    _configurar_tabela(tabela, bloco.largura_total(), bordas)
    _grade(tabela, [c.largura for c in bloco.colunas])

    def escrever(valores, primeira, cabecalho=False, negrito=False):
        celulas = tabela.rows[0].cells if primeira else tabela.add_row().cells
        linha = tabela.rows[-1]
        propriedades = linha._tr.get_or_add_trPr()
        _filho(propriedades, "w:cantSplit")                 # a linha não se parte entre duas páginas
        if cabecalho:
            _filho(propriedades, "w:tblHeader")             # o cabeçalho se repete em cada página
            if bloco.altura_cabecalho:
                linha.height, linha.height_rule = Twips(bloco.altura_cabecalho), WD_ROW_HEIGHT_RULE.AT_LEAST
        for coluna, (celula, valor) in enumerate(zip(celulas, valores)):
            celula.width = Twips(bloco.colunas[coluna].largura)
            celula.vertical_alignment = WD_ALIGN_VERTICAL.CENTER if bloco.altura_cabecalho and cabecalho else None
            p = celula.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cabecalho else ALINHAMENTO_DOCX[bloco.colunas[coluna].alinhamento]
            _fonte(p.add_run(_quebras(str(valor))), doc.fonte, bloco.tamanho, negrito=cabecalho or negrito)
            if cabecalho and bloco.sombreado:
                _filho(celula._tc.get_or_add_tcPr(), "w:shd", val="clear", color="auto", fill="D9D9D9")

    primeira = True
    if bloco.cabecalho:
        escrever([c.titulo for c in bloco.colunas], True, cabecalho=True)
        primeira = False
    for posicao, valores in enumerate(bloco.linhas):
        escrever(valores, primeira, negrito=posicao in bloco.negrito_linhas)
        primeira = False


def _formulario_docx(d, doc, bloco):
    tabela = d.add_table(rows=len(bloco.linhas), cols=len(bloco.grade))
    _configurar_tabela(tabela, bloco.largura_total(), {lado: "000000" for lado in ("top", "left", "bottom", "right", "insideH", "insideV")},
                       margens=(0, 108, 0, 108))
    _grade(tabela, bloco.grade)                              # antes de mesclar: a célula mesclada soma as larguras
    for r, (linha, altura) in enumerate(zip(bloco.linhas, bloco.alturas)):
        tabela.rows[r].height, tabela.rows[r].height_rule = Twips(altura), WD_ROW_HEIGHT_RULE.AT_LEAST
        _filho(tabela.rows[r]._tr.get_or_add_trPr(), "w:cantSplit")
        coluna = 0
        for celula in linha:
            inicio = tabela.cell(r, coluna)
            alvo = inicio.merge(tabela.cell(r, coluna + celula.span - 1)) if celula.span > 1 else inicio
            p = alvo.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            if celula.texto:
                _fonte(p.add_run(_quebras(celula.texto)), doc.fonte, bloco.tamanho, negrito=celula.negrito)
            coluna += celula.span


def _pauta_docx(d, doc, bloco, largura):
    tabela = d.add_table(rows=bloco.linhas, cols=1)
    _configurar_tabela(tabela, largura, {"bottom": "000000", "insideH": "000000"}, margens=(0, 0, 0, 0))
    _grade(tabela, [largura])
    for linha in tabela.rows:
        linha.height, linha.height_rule = Twips(int(bloco.altura * 20)), WD_ROW_HEIGHT_RULE.EXACTLY
        p = linha.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _fonte(p.add_run(""), doc.fonte, 6)


def _imagem_docx(d, doc, bloco):
    p = d.add_paragraph()
    p.alignment = ALINHAMENTO_DOCX[bloco.alinhamento]
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(bloco.antes), Pt(0)
    p.add_run().add_picture(str(bloco.caminho), width=Mm(bloco.largura))


def _colunas_docx(d, doc, bloco, largura):
    """Tabela sem bordas: uma coluna de verdade por coluna do bloco, com uma coluna vazia no meio para o espaço."""
    n = len(bloco.colunas)
    util = (largura - bloco.espaco * (n - 1)) // n
    larguras = [x for i in range(n) for x in ((util,) if i == 0 else (bloco.espaco, util))]
    tabela = d.add_table(rows=1, cols=len(larguras))
    _configurar_tabela(tabela, sum(larguras), None, margens=(0, 0, 0, 0))
    _grade(tabela, larguras)
    for celula, paragrafos in zip(tabela.rows[0].cells[::2], bloco.colunas):
        for paragrafo in paragrafos:
            _paragrafo_docx(d, doc, paragrafo, alvo=celula, quebra=False)
        inicial = celula.paragraphs[0]                              # o parágrafo vazio que toda célula já traz
        if not inicial.text and len(celula.paragraphs) > 1:
            inicial._element.getparent().remove(inicial._element)


def _assinaturas_docx(d, doc, bloco, largura):
    n = len(bloco.rotulos)
    larguras = [largura // n] * n
    tabela = d.add_table(rows=2, cols=n)
    _configurar_tabela(tabela, sum(larguras), None, margens=(0, 0, 0, 0))
    _grade(tabela, larguras)
    for coluna, rotulo in enumerate(bloco.rotulos):
        for linha, (texto, tamanho) in enumerate((("_" * bloco.comprimento, doc.tamanho), (rotulo, bloco.tamanho_rotulo))):
            p = tabela.cell(linha, coluna).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            _fonte(p.add_run(texto), doc.fonte, tamanho)


def _assinatura_docx(d, doc, bloco):
    texto = lambda conteudo, antes=0, **fonte: _paragrafo_docx(d, doc, Paragrafo([Texto(conteudo, **fonte)], "centro", antes=antes, unido_ao_proximo=True))
    texto("_" * 44, antes=bloco.antes)
    texto(bloco.nome, negrito=True)
    texto(bloco.cargo)
    if bloco.registro:
        texto(bloco.registro)
    texto(bloco.papel, italico=True).paragraph_format.keep_with_next = False


def para_docx(doc):
    """Documento -> BytesIO com o arquivo .docx pronto para enviar."""
    d = _base_docx(doc)
    for numero, folha in enumerate(doc.paginas()):
        if numero == 0:
            _configurar_secao(d.sections[0], doc, folha)
        elif folha.secao:                                    # página diferente (paisagem, sem timbre): seção nova
            secao = d.add_section(WD_SECTION.NEW_PAGE)
            _miniatura(d.paragraphs[-1])                     # o parágrafo com a marca de fim da seção anterior não pode abrir uma página em branco
            _configurar_secao(secao, doc, folha)
        largura = doc.largura_da(folha)
        for posicao, bloco in enumerate(folha.blocos):
            # o primeiro bloco de uma página que já começa em página nova não precisa de outra quebra
            quebra = getattr(bloco, "quebra_antes", False) and not (posicao == 0 and (numero == 0 or folha.secao))
            if bloco.tipo == "titulo":
                _paragrafo_docx(d, doc, Paragrafo([Texto(bloco.texto, negrito=True)], "centro", antes=bloco.antes,
                                                  depois=bloco.depois, tamanho=bloco.tamanho))
            elif bloco.tipo == "paragrafo":
                _paragrafo_docx(d, doc, bloco, quebra=quebra)
            elif bloco.tipo in ("tabela", "formulario", "pauta", "assinaturas", "colunas"):
                if bloco.antes:                              # o espaço antes de uma tabela é um parágrafo vazio
                    _paragrafo_docx(d, doc, Paragrafo([""], antes=bloco.antes, tamanho=1, unido_ao_proximo=True))      # junto da tabela: a pergunta que vem antes não se solta das linhas dela
                if bloco.tipo in ("tabela", "formulario"):
                    {"tabela": _tabela_docx, "formulario": _formulario_docx}[bloco.tipo](d, doc, bloco)
                else:
                    {"pauta": _pauta_docx, "assinaturas": _assinaturas_docx, "colunas": _colunas_docx}[bloco.tipo](d, doc, bloco, largura)
            elif bloco.tipo == "imagem":
                _imagem_docx(d, doc, bloco)
            elif bloco.tipo == "assinatura":
                _assinatura_docx(d, doc, bloco)
    propriedades = d.core_properties
    propriedades.title, propriedades.author, propriedades.last_modified_by = doc.titulo, "Sistema de Prontuários", "Sistema de Prontuários"
    propriedades.comments = propriedades.subject = propriedades.keywords = ""
    arquivo = BytesIO()
    d.save(arquivo)
    arquivo.seek(0)
    return arquivo
